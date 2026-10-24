from pathlib import Path
import re,json,hashlib,zipfile
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START, WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pypdf import PdfReader
R=Path(__file__).resolve().parents[1]; ROOT=R.parents[1]
data=json.loads((R/'evidence/audit-results-current.json').read_text(encoding='utf-8'))
manifest=json.loads((R/'evidence/screenshots-manifest.json').read_text(encoding='utf-8'))
shots={Path(x['file']).stem:R/x['file'] for x in manifest}
D=Document();S=D.sections[0];S.page_width=Cm(21);S.page_height=Cm(29.7)
S.top_margin=Cm(2.4);S.bottom_margin=Cm(2.2);S.left_margin=Cm(2.6);S.right_margin=Cm(2.1)
S.header_distance=Cm(1.0);S.footer_distance=Cm(1.0);S.different_first_page_header_footer=True
for name in ['Normal','Body Text','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption','TOC 1','TOC 2']:
    try:st=D.styles[name]
    except KeyError:continue
    st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
    st._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Times New Roman')
    st.font.size=Pt(13);st.paragraph_format.space_after=Pt(5);st.paragraph_format.line_spacing=1.15
normal=D.styles['Normal'];normal.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
normal.paragraph_format.widow_control=True
for name in ['TOC 1','TOC 2','Table of Figures']:
    if name not in D.styles:continue
    D.styles[name].font.size=Pt(11)
    D.styles[name].paragraph_format.line_spacing=1.0
    D.styles[name].paragraph_format.space_after=Pt(2)
for name,sz in [('Title',22),('Heading 1',16),('Heading 2',14),('Heading 3',13),('Caption',11)]:
    D.styles[name].font.size=Pt(sz)
for name in ['Heading 1','Heading 2','Heading 3']:
    D.styles[name].font.bold=True;D.styles[name].paragraph_format.keep_with_next=True
    D.styles[name].paragraph_format.space_before=Pt(9)
D.styles['Heading 1'].paragraph_format.page_break_before=True
D.styles['Caption'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
D.styles['Caption'].paragraph_format.space_before=Pt(3)
settings=D.settings.element;u=OxmlElement('w:updateFields');u.set(qn('w:val'),'true');settings.append(u)
idn=10;figcount=0;tabcount=0;section_title='';chapter=0;sub=0

def field(p,code,result=''):
    b=OxmlElement('w:fldChar');b.set(qn('w:fldCharType'),'begin');p.add_run()._r.append(b)
    t=OxmlElement('w:instrText');t.set(qn('xml:space'),'preserve');t.text=' '+code+' ';p.add_run()._r.append(t)
    sep=OxmlElement('w:fldChar');sep.set(qn('w:fldCharType'),'separate');p.add_run()._r.append(sep)
    if result:p.add_run(result)
    e=OxmlElement('w:fldChar');e.set(qn('w:fldCharType'),'end');p.add_run()._r.append(e)
def bookmark(p,name):
    global idn
    idn+=1;b=OxmlElement('w:bookmarkStart');b.set(qn('w:id'),str(idn));b.set(qn('w:name'),name);p._p.insert(0,b)
    e=OxmlElement('w:bookmarkEnd');e.set(qn('w:id'),str(idn));p._p.append(e)
def para(text='',style=None):
    p=D.add_paragraph(style=style)
    for i,part in enumerate(re.split(r'(\[\[fig:[^]]+\]\])',text)):
        if part.startswith('[[fig:'):field(p,'REF fig_'+part[6:-2].replace('-','_')+' \\h','?')
        else:p.add_run(part)
    return p
def caption(kind,name,title):
    global figcount,tabcount
    if kind=='Hình':figcount+=1;count=figcount;seq='Figure';prefix='fig_'
    else:tabcount+=1;count=tabcount;seq='Table';prefix='tab_'
    p=D.add_paragraph(style='Caption');p.add_run(kind+' ')
    # Bookmark the SEQ only so REF gives the number rather than the entire caption.
    idname=prefix+name.replace('-','_');start=OxmlElement('w:bookmarkStart');global idn;idn+=1;start.set(qn('w:id'),str(idn));start.set(qn('w:name'),idname);p._p.append(start)
    field(p,'SEQ '+seq+' \\* ARABIC',str(count));end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(idn));p._p.append(end)
    p.add_run('  '+title);return p
def table(headers,rows,title=None):
    name='table_'+str(tabcount+1)
    p=D.add_paragraph();p.paragraph_format.space_after=Pt(3);p.add_run('Bảng ');field(p,'REF tab_'+name+' \\h',str(tabcount+1));p.add_run(' tổng hợp '+(title or section_title).lower()+'.')
    cap=caption('Bảng',name,title or section_title);cap.paragraph_format.keep_with_next=True
    t=D.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    widths=[16.3/len(headers)]*len(headers)
    if len(headers)==3:widths=[4.0,5.5,6.8]
    if len(headers)==4:widths=[3.0,4.1,4.6,4.6]
    if len(headers)==5:widths=[3.2,3.4,3.4,3.5,2.8]
    for c,w in zip(t.columns,widths):c.width=Cm(w)
    def fill(c,txt,header=False,idx=0):
        c.text=txt;c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        tcPr=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
        for edge in ['top','left','bottom','right']:
            z=OxmlElement('w:'+edge);z.set(qn('w:val'),'single');z.set(qn('w:sz'),'4');z.set(qn('w:color'),'D9D9D9');b.append(z)
        tcPr.append(b);m=OxmlElement('w:tcMar')
        for edge in ['top','left','bottom','right']:
            z=OxmlElement('w:'+edge);z.set(qn('w:w'),'85');z.set(qn('w:type'),'dxa');m.append(z)
        tcPr.append(m)
        sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E7ECEF' if header else 'F7F8F9' if idx%2 else 'FFFFFF');tcPr.append(sh)
        for p in c.paragraphs:
            p.alignment=WD_ALIGN_PARAGRAPH.LEFT;p.paragraph_format.line_spacing=1.05;p.paragraph_format.space_after=Pt(2)
            for r in p.runs:r.font.size=Pt(11);r.bold=header
    for c,h,w in zip(t.rows[0].cells,headers,widths):c.width=Cm(w);fill(c,h,True)
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for idx,row in enumerate(rows):
        tr=t.add_row()
        no_split=OxmlElement('w:cantSplit');tr._tr.get_or_add_trPr().append(no_split)
        for c,txt,w in zip(tr.cells,row,widths):c.width=Cm(w);fill(c,str(txt),False,idx)
    para()
def diagram(name,title):
    from PIL import Image
    section=D.add_section(WD_SECTION_START.NEW_PAGE)
    section.orientation=WD_ORIENT.LANDSCAPE;section.page_width=Cm(29.7);section.page_height=Cm(21)
    section.different_first_page_header_footer=False
    section.top_margin=Cm(1.2);section.bottom_margin=Cm(1.2);section.left_margin=Cm(2.0);section.right_margin=Cm(2.0)
    p=D.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
    file=R/'assets/diagrams'/f'{name}.png';w,h=Image.open(file).size
    width=min(24.8,16.1*w/h)
    p.add_run().add_picture(str(file),width=Cm(width))
    caption('Hình',name,title)
    if name=='fund-flow':
        p=para('Trong harness local, 72.000đ donation demo kết thúc thành 36.000đ available, 0đ reserved và 36.000đ spent. Clock được tiêm; HEAD ảnh được stub. Không có giao hàng, upload hoặc chuyển khoản thật. Mỗi phase lưu JSON/ledger và assert cân bằng, không gán thẳng status hoặc số dư.')
        p.paragraph_format.line_spacing=1.0
        for rr in p.runs:rr.font.size=Pt(10)
    section=D.add_section(WD_SECTION_START.NEW_PAGE)
    section.orientation=WD_ORIENT.PORTRAIT;section.page_width=Cm(21);section.page_height=Cm(29.7)
    section.different_first_page_header_footer=False
    section.top_margin=Cm(2.4);section.bottom_margin=Cm(2.2);section.left_margin=Cm(2.6);section.right_margin=Cm(2.1)
def pictures(names,title):
    name='ui_'+names[0];p=D.add_paragraph();p.add_run('Hình ');field(p,'REF fig_'+name.replace('-','_')+' \\h',str(figcount+1));p.add_run(' ghi lại giao diện đang chạy trong môi trường demo.')
    p.paragraph_format.keep_with_next=True
    t=D.add_table(rows=1,cols=len(names));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for i,(cell,n) in enumerate(zip(t.rows[0].cells,names)):
        cell.width=Cm(8.1);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
        p=cell.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        from PIL import Image
        w,h=Image.open(shots[n]).size
        width=6.15 if h>w else 8.0
        p.add_run().add_picture(str(shots[n]),width=Cm(width))
        p.paragraph_format.space_after=Pt(0);p.paragraph_format.keep_with_next=True
        q=cell.add_paragraph('('+chr(97+i)+') '+n);q.alignment=WD_ALIGN_PARAGRAPH.CENTER;q.paragraph_format.keep_with_next=True
        for rr in q.runs:rr.font.size=Pt(9)
    caption('Hình',name,title)

header=S.header.paragraphs[0];header.text='RescueFarm   |   Đồ án Phát triển ứng dụng Mobile'
header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
for rr in header.runs:rr.font.size=Pt(9)
footer=S.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER;field(footer,'PAGE','1')
D.core_properties.title='RescueFarm ứng dụng Android giải cứu nông sản';D.core_properties.author='Nhóm RescueFarm';D.core_properties.subject='Báo cáo từ mã nguồn và bằng chứng kiểm thử local'

# Cover follows supplied academic hierarchy; title and metadata remain editable.
for txt in ['BỘ GIÁO DỤC VÀ ĐÀO TẠO','ĐẠI HỌC KINH TẾ TP HỒ CHÍ MINH UEH','TRƯỜNG CÔNG NGHỆ VÀ THIẾT KẾ']:
    p=para(txt);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for rr in p.runs:rr.bold=True;rr.font.size=Pt(14)
p=para();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run().add_picture(str(R/'assets/diagrams/template-X10.png'),width=Cm(9))
p=para('ĐỒ ÁN MÔN HỌC','Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p=para('XÂY DỰNG ỨNG DỤNG BÁN NÔNG SẢN GIẢI CỨU','Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p=para('RescueFarm');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
for txt in ['Học phần  Phát triển ứng dụng Mobile','Chuyên ngành  Công nghệ thông tin   |   Khóa K50','Giảng viên  TS. Đặng Ngọc Hoàng Thành','Nhóm 7']:
    p=para(txt);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
for txt in ['Lại Thành Đạt','Phan Khắc Anh Tuấn','Nguyễn Phương Chinh','Nguyễn Tấn Khiêm','Trương Thành Đạt']:
    p=para(txt);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(2)
p=para('Bản nháp chờ xác nhận MSSV lớp và phân công');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p=para('Audit mã nguồn ngày 24 tháng 10 năm 2026');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
D.add_page_break()

para('PHẠM VI VÀ CÁCH ĐỌC BẰNG CHỨNG','Heading 1').paragraph_format.page_break_before=False
para('Báo cáo phân tích RescueFarm như một hệ thống thương mại điện tử giải cứu nông sản. Phần chính giải thích quan hệ giữa nghiệp vụ, giao diện Android, dữ liệu và backend; ảnh và kết quả kiểm thử được thu thập từ môi trường demo cô lập. Chương 1 giữ nội dung do nhóm cung cấp trong PDF; các chương sau được tổng hợp từ repository hiện tại.')
para('HEAD nền  '+data['head']+'. Working tree có thay đổi phục vụ thực nghiệm chưa commit và thay đổi prompt có sẵn được giữ nguyên. APK demo có package riêng và SHA-256 trong evidence. Không tự commit, push hoặc deploy; không sử dụng dữ liệu hay giao dịch production để tạo ảnh.')
para('Quy ước đường dẫn: ui/, domain/, data/ và service/ được rút gọn từ app/src/main/java/com/rescuefarm/; layout và navigation thuộc app/src/main/res/; functions/ và cloudinary-worker/ tính từ root repository. Chương 1 giữ nguyên nội dung nhóm; nhận định về thiết bị vật lý trong bản nhóm chưa có bằng chứng ở lượt audit này, vốn chỉ dùng emulator.')
table(['Nhãn bằng chứng','Cách hiểu'],[
['Có code','Đã đọc đường triển khai; chưa đồng nghĩa runtime hoặc provider ngoài đã hoạt động.'],
['Đã qua test','Bài kiểm thử đã chạy trong lượt audit, với phạm vi và log ghi rõ.'],
['Đã thao tác trên emulator','Activity/Fragment thật đã chạy với dữ liệu demo; không suy rộng mọi nút.'],
['Chưa xác minh tích hợp ngoài','OAuth, ngân hàng, upload, FCM hoặc deploy live chưa được thực nghiệm.'],
['Thiếu hoặc chưa hoàn thiện','Thiếu thông tin nhóm, ca demo hay yêu cầu học thuật được nêu cụ thể.']],title='Các mức bằng chứng sử dụng trong báo cáo')
para('MSSV, lớp cụ thể, nhiệm vụ được giao và tỷ lệ hoàn thành chờ nhóm xác nhận. Chưa có kết quả kiểm tra tương đồng UEH ở chế độ mặc định; bản báo cáo không cam kết dưới 20% khi chưa có minh chứng thật.')

para('MỤC LỤC','Heading 1');p=para();field(p,'TOC \\o "1-2" \\h \\z \\u','Mục lục sẽ được cập nhật khi render bằng Word.')
para('DANH MỤC HÌNH VÀ BẢNG','Heading 1');p=para();field(p,'TOC \\c "Figure" \\h \\z','Danh mục hình tự động');p=para();field(p,'TOC \\c "Table" \\h \\z','Danh mục bảng tự động')

# Preserve original Chapter 1 from supplied PDF. Normalize extraction whitespace only.
original=' '.join(PdfReader(ROOT/'docs/MOBILE FINAL PROJECT.pdf').pages[i].extract_text() for i in [3,4,5]);original=re.sub(r'\s+',' ',original).strip()
headings=['CHƯƠNG 1. TỔNG QUAN','1.1. Giới thiệu','1.2. Tính Cấp Thiết của Đề tài','1.3. Công Cụ Phát Triển Ứng Dụng Di động','a) Giới thiệu sơ lược về Java','b) Android SDK','c) Android Studio','d) Firebase','1.4. Kiểm thử ứng dụng trên Android Emulator và Thiết bị']
pattern='|'.join(re.escape(x) for x in headings);parts=re.split('('+pattern+')',original)
for part in parts:
    part=part.strip()
    if not part:continue
    if part in headings:
        para(part.replace('CHƯƠNG 1.','CHƯƠNG 1'), 'Heading 1' if part.startswith('CHƯƠNG') else 'Heading 3' if re.match('[a-d]\\)',part) else 'Heading 2')
    else:
        for chunk in re.split(r'(?=RescueFarm là |Hệ thống phục vụ |Một khó khăn |Về phương diện |Bên cạnh Firebase|Thiết bị Android|Ngoài kiểm tra)',part):
            if chunk.strip():para(chunk.strip())
# The verification note belongs to the evidence scope page; keep Chapter 1 unchanged.

compact={'Bài toán và ranh giới hệ thống','Tác nhân và quyền sở hữu','Yêu cầu đơn hàng và giao nhận','Yêu cầu giải cứu mở rộng','Danh mục dữ liệu người dùng và hàng hóa','Chiến dịch và phân bổ theo lô','Dữ liệu giải cứu và các collection vận hành','TTL và đồng bộ cục bộ','Xác thực và Google Credential Manager','Triển khai Home và tìm kiếm','Triển khai tương tác và quản trị','Đồng bộ và xử lý lỗi','Quy trình tài khoản và thêm sản phẩm','Rủi ro còn lại và kết luận kỹ thuật','Bảng phân công chờ xác nhận của nhóm'}
text=(R/'source/report-content.md').read_text(encoding='utf-8').replace('firebase.schema.json','firestore.schema.json')
text=re.sub(r'Mỗi bước lưu JSON trạng thái và ledger;.*?không chứng minh có tiền thật hoặc bàn giao vật lý\.\n','',text,flags=re.S)
for name in ['architecture','oop','data-model','order-flow','fund-flow']:
    pattern=r'(## [^\n]+\n)(.*?)(?=\n##? |\Z)'
    def move_diagram(m):
        body=m.group(2);fig=re.search(r'\{\{FIG:'+re.escape(name)+r'\|[^\n]+\}\}',body)
        if not fig:return m.group(0)
        return m.group(1)+body.replace(fig.group(0),'').rstrip()+'\n\n'+fig.group(0)+'\n'
    text=re.sub(pattern,move_diagram,text,flags=re.S)
lines=text.splitlines();i=0
while i<len(lines):
    line=lines[i].strip()
    if not line:i+=1;continue
    if line.startswith('# '):
        title=line[2:];m=re.match(r'CHƯƠNG (\d+)',title);chapter=int(m.group(1)) if m else 0;sub=0
        para(title,'Heading 1');i+=1;continue
    if line.startswith('## '):
        section_title=line[3:];sub+=1;prefix=f'{chapter}.{sub}  ' if chapter else ''
        p=para(prefix+section_title,'Heading 2');p.paragraph_format.page_break_before=(chapter==3 and sub>1)
        # Never create a blank chapter title page: first subsection stays with chapter title.
        if sub==1:p.paragraph_format.page_break_before=False
        i+=1;continue
    if line.startswith('{{FIG:'):
        n,title=line[6:-2].split('|',1);diagram(n,title);i+=1;continue
    if line.startswith('{{SHOTS:'):
        n,title=line[8:-2].split('|',1);pictures(n.split(','),title);i+=1;continue
    if line=='{{RESULTS}}':
        j=data['jvm'];li=data['lint'];table(['Kiểm tra','Kết quả hiện tại','Giới hạn'],[
        ['JVM Android',f"{j['tests']} test; {j['failures']} failure; {j['errors']} error; {j['skipped']} skipped",'Không phải số chức năng hoặc coverage'],
        ['assembleDebug + lintDebug','Build qua; '+str(li['Error']+li['Fatal'])+' lỗi; '+str(li['Warning'])+' cảnh báo Lint','Cảnh báo còn cần phân loại'],
        ['Backend check và unit','Exit 0; 183 qua, 0 lỗi, 0 skipped','Logic local, không provider live'],
        ['Cloudinary Worker','9 qua, 0 lỗi, 0 skipped','Không upload thật'],
        ['Firestore/API Emulator','173 qua, 0 lỗi, 0 skipped','Project demo-rescuefarm; bao gồm Hũ'],
        ['Navigation','59 destination Fragment có class','Kiểm tra tĩnh'],
        ['Instrumentation chọn lọc','11 qua, 0 lỗi, 0 skipped','9 layout + 1 Room + 1 rotation'],
        ['Ảnh UI và Room restart','Test chụp actor và 2 method offline qua','Ảnh thật, phạm vi theo manifest']],title='Kết quả kiểm thử mới trên revision được audit');i+=1;continue
    if line=='{{GITTABLE}}':
        g=json.loads((R/'evidence/git-evidence.json').read_text(encoding='utf-8'))
        rows=[[x['sha'][:7],x['subject'],', '.join(x.get('files',[])[:3]) or 'Tra cứu diff merge trong evidence']for x in g['merges'][:6]]
        table(['Merge SHA','Branch hoặc PR trong lịch sử','Sản phẩm tra cứu'],rows,title='Các merge gần HEAD để kiểm tra tích hợp');i+=1;continue
    if line=='{{TEAMTABLE}}':
        table(['Thành viên','Nhiệm vụ và bàn giao','Module và merge','Trạng thái và phần trăm'],[[name,'Chờ nhóm xác nhận nhiệm vụ được giao và sản phẩm đã duyệt','Tra git-evidence; alias và sở hữu module chờ xác nhận','Chờ xác nhận']for name in ['Lại Thành Đạt','Phan Khắc Anh Tuấn','Nguyễn Phương Chinh','Nguyễn Tấn Khiêm','Trương Thành Đạt']],title='Phân công cần nhóm xác nhận trước nộp');i+=1;continue
    if line.startswith('|'):
        block=[]
        while i<len(lines) and lines[i].strip().startswith('|'):block.append([c.strip() for c in lines[i].strip().strip('|').split('|')]);i+=1
        table(block[0],[r for r in block[1:] if not all(re.fullmatch(r'[-:]+',c) for c in r)]);continue
    p=para(line)
    if line.startswith('['):
        p.paragraph_format.line_spacing=1.1
        for rr in p.runs:rr.font.size=Pt(12)
    i+=1

# Remove inherited Office title rules; a title is distinguished by typography.
for element in list(D.element.iter(qn('w:pBdr')))+list(D.styles.element.iter(qn('w:pBdr'))):
    element.getparent().remove(element)
D.save(R/'report.docx')
(R/'source/build-stats.json').write_text(json.dumps(dict(figures=figcount,tables=tabcount,images=len(D.inline_shapes),sourceWords=len(text.split())),indent=2),encoding='utf-8')
print('DOCX written',figcount,'figures',tabcount,'tables',len(D.inline_shapes),'embedded images')
