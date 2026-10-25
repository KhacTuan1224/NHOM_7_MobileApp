"""Build only chapter 5. Run with the Codex bundled Python runtime."""
from pathlib import Path
import html, json, math, shutil, hashlib, subprocess
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT=Path(__file__).resolve().parents[1]; ROOT=OUT.parents[2]
AS=OUT/'assets'; AS.mkdir(exist_ok=True)
FONT='C:/Windows/Fonts/arial.ttf'; BOLD='C:/Windows/Fonts/arialbd.ttf'

def sequence(name, actors, steps, note=''):
    """Editable SVG plus 2x PNG; dashed messages are responses, not DB relations."""
    w=1320; h=110+len(steps)*55+48; scale=2
    im=Image.new('RGB',(w*scale,h*scale),'white'); dr=ImageDraw.Draw(im)
    svg=[f'<svg xmlns="http://www.w3.org/2000/svg" width="{w}" height="{h}" viewBox="0 0 {w} {h}"><rect width="100%" height="100%" fill="white"/>']
    xs=[145+i*(w-290)/(len(actors)-1) for i in range(len(actors))]
    def text(x,y,s,size=23,bold=False):
        ft=ImageFont.truetype(BOLD if bold else FONT,size*scale)
        dr.text((x*scale,y*scale),s,font=ft,fill='#172b31',anchor='mt')
        svg.append(f'<text x="{x}" y="{y+size}" text-anchor="middle" font-family="Arial" font-size="{size}" font-weight="{700 if bold else 400}" fill="#172b31">{html.escape(s)}</text>')
    def line(x1,y1,x2,y2,color='#7d8e95',dash=False,arrow=False):
        if dash:
            length=math.hypot(x2-x1,y2-y1)
            for d in range(0,int(length),14):
                a=d/length;b=min(d+7,length)/length
                dr.line(((x1+(x2-x1)*a)*scale,(y1+(y2-y1)*a)*scale,(x1+(x2-x1)*b)*scale,(y1+(y2-y1)*b)*scale),fill=color,width=2*scale)
        else:dr.line((x1*scale,y1*scale,x2*scale,y2*scale),fill=color,width=2*scale)
        svg.append(f'<path d="M{x1} {y1} L{x2} {y2}" stroke="{color}" stroke-width="2" fill="none"'+(' stroke-dasharray="7 7"' if dash else '')+'/>')
        if arrow:
            sign=1 if x2>x1 else -1
            pts=[(x2,y2),(x2-sign*12,y2-6),(x2-sign*12,y2+6)]
            dr.polygon([(x*scale,y*scale) for x,y in pts],fill=color)
            svg.append('<polygon points="'+' '.join(f'{x},{y}' for x,y in pts)+f'" fill="{color}"/>')
    for x,a in zip(xs,actors):
        # Headers may wrap; preserve short column labels at publication size.
        bw=min(252,(w-70)/len(actors)-15)
        dr.rounded_rectangle(((x-bw/2)*scale,7*scale,(x+bw/2)*scale,65*scale),radius=8*scale,fill='#eaf2ee',outline='#35664b',width=2*scale)
        svg.append(f'<rect x="{x-bw/2}" y="7" width="{bw}" height="58" rx="8" fill="#eaf2ee" stroke="#35664b" stroke-width="2"/>')
        lines=a.split('\n')
        for i,s in enumerate(lines):text(x,15+i*24,s,21,True)
        line(x,70,x,h-53,dash=True)
    for i,(a,b,label,*opt) in enumerate(steps):
        y=110+i*55; x1=xs[a];x2=xs[b]; dash=bool(opt and opt[0])
        ft=ImageFont.truetype(FONT,23*scale); tw=dr.textlength(label,font=ft)/scale+16; cx=(x1+x2)/2
        line(x1,y,x2,y,'#315943',dash,True)
        dr.rectangle(((cx-tw/2)*scale,(y-32)*scale,(cx+tw/2)*scale,(y-5)*scale),fill='white')
        svg.append(f'<rect x="{cx-tw/2}" y="{y-32}" width="{tw}" height="27" fill="white"/>');text(cx,y-32,label)
    if note:text(w/2,h-34,note,20)
    svg.append('</svg>'); (AS/(name+'.svg')).write_text('\n'.join(svg),encoding='utf-8');im.save(AS/(name+'.png'))
    (AS/(name+'.json')).write_text(json.dumps(dict(actors=actors,steps=steps,note=note),ensure_ascii=False,indent=2),encoding='utf-8')

sequence('firebase-connection',['Android\nJava XML','Firebase\nAuth và Rules','Cloud\nFirestore','Render\nAdmin SDK'],[
 (0,1,'1  Đăng nhập email hoặc Google credential'),(1,0,'2  Phiên người dùng và ID token',True),
 (0,1,'3  SDK đọc ghi theo UID và vai trò'),(1,2,'4  Rules cho phép hoặc từ chối'),
 (0,3,'5  HTTP kèm ID token cho nghiệp vụ tin cậy'),(3,2,'6  Xác thực và transaction phía server'),
 (2,0,'7  Snapshot cập nhật các màn đang theo dõi',True)],'Rules kiểm soát SDK client; backend Admin SDK phải tự kiểm tra quyền.')
sequence('home-search',['Home và\ntìm kiếm','ViewModel\nDiscoveryEngine','Repository\nRoom cache','Firestore\nSDK'],[
 (0,1,'1  Mở trang hoặc đổi từ khóa bộ lọc'),(1,2,'2  Theo dõi danh mục và dữ liệu công khai'),
 (2,1,'3  Trả bản cache đủ điều kiện nếu có',True),(2,3,'4  Query ACTIVE có giới hạn trang'),
 (3,2,'5  Snapshot mới và ánh xạ model',True),(2,1,'6  Lưu Room và phát LiveData',True),
 (1,0,'7  Lọc sắp xếp và dựng các nhóm nội dung',True)],'Đổi bộ lọc dùng dữ liệu đã tải; tải thêm mới mở rộng tập kết quả.')
sequence('commerce',['Chi tiết\ngiỏ và checkout','Repository\nAndroid','Render\ncommerce.js','Firestore\ntransaction'],[
 (0,1,'1  Chọn productId batchId và số lượng'),(1,2,'2  Lấy báo giá hiện hành'),
 (1,0,'3  Giỏ Guest ở Room hoặc Customer ở carts',True),(0,1,'4  Xác nhận checkout với requestId'),
 (1,2,'5  checkoutGuest hoặc checkoutCustomer'),(2,3,'6  Kiểm quyền giá hạn lô và tồn khả dụng'),
 (3,2,'7  Ghi đơn giữ kho và receipt cùng transaction',True),(2,0,'8  Trả đơn hoặc yêu cầu xác nhận lại báo giá',True)],'QR chỉ được ghi nhận thanh toán qua webhook SePay hợp lệ; COD theo vòng đời đơn.')
sequence('campaign-map',['Seller hoặc\nCustomer','Android\nRepository','Firestore\ncampaigns','Render và\nmedia'],[
 (0,1,'1  Seller bật chia sẻ GPS của chiến dịch'),(1,2,'2  Ghi fix mới timestamp và cờ chia sẻ'),
 (2,3,'3  Listener chung nhận thay đổi',True),(3,1,'4  SSE rescue-map/events gửi điểm công khai',True),
 (1,0,'5  Cập nhật marker và thời điểm vị trí',True),(0,3,'6  Customer tải ảnh qua Worker và Cloudinary'),
 (0,3,'7  Gửi text ảnh và lựa chọn ẩn danh'),(3,2,'8  Xác thực tác giả rồi ghi comments transaction'),
 (2,1,'9  Listener bình luận cập nhật danh sách',True)],'GPS có chu kỳ; SSE truyền thay đổi đã ghi, không tự đo vị trí xe.')
sequence('seller-admin',['Seller\ntrên mobile','Repository\nvà media','Firestore\nvà Rules','Admin\ntrên mobile'],[
 (0,1,'1  Nhập sản phẩm lô hoặc chiến dịch'),(1,2,'2  Kiểm owner và ghi dữ liệu hợp lệ'),
 (2,3,'3  Nội dung chờ duyệt xuất hiện',True),(3,2,'4  Duyệt hoặc từ chối theo quyền Admin'),
 (2,1,'5  Snapshot trạng thái mới',True),(1,0,'6  Seller xem kết quả và cập nhật hợp lệ',True)],'Chức năng quản trị nằm trong ứng dụng Android; không có web admin được suy diễn.')
sequence('rescue-fund',['Người ủng hộ\nSeller và Admin','Render\nrescue-fund.js','Firestore\nquỹ kho và ledger','SePay hoặc\nngân hàng'],[
 (0,3,'1  Ủng hộ theo QR và mã theo dõi'),(3,1,'2  Webhook tiền vào hợp lệ'),
 (1,2,'3  Chống trùng và ghi nhận PAID'),(1,2,'4  Phân bổ giữ đồng thời tiền và hàng'),
 (0,1,'5  Seller ACCEPT rồi DELIVER kèm ảnh trước hạn'),(1,2,'6  Ghi bằng chứng hoặc hủy quá hạn giải phóng'),
 (0,1,'7  Admin VERIFY chuyến đã giao hợp lệ'),(0,3,'8  Admin chuyển tiền ngoài ứng dụng'),
 (0,1,'9  PAY ghi mã giao dịch số tiền và chứng từ'),(1,2,'10  Chuyển reserved sang spent và ghi ledger')],'DELIVER đúng hạn được duyệt muộn; RESERVED hoặc ACCEPTED quá hạn mới tự hủy.')

shutil.copyfile(ROOT/'docs/diagram (2).png',AS/'architecture-gitdiagram-hd.png')
with Image.open(AS/'architecture-gitdiagram-hd.png') as original:
    original.thumbnail((3000,3000))
    original.save(AS/'architecture-gitdiagram-print.png')
D=Document(); sec=D.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(1.8);sec.bottom_margin=Cm(1.8);sec.left_margin=Cm(2.3);sec.right_margin=Cm(2.0)
sec.header_distance=Cm(.75);sec.footer_distance=Cm(.8)
for name in ['Normal','Title','Heading 1','Heading 2','Heading 3','Caption']:
    s=D.styles[name];s.font.name='Times New Roman';s.font.color.rgb=RGBColor(0,0,0)
    s._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Times New Roman')
    s.font.size=Pt(12.5);s.paragraph_format.space_after=Pt(5);s.paragraph_format.line_spacing=1.08
D.styles['Normal'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
for name,size in [('Title',16),('Heading 1',14),('Heading 2',13),('Heading 3',12.5)]:
    s=D.styles[name];s.font.size=Pt(size);s.font.bold=True;s.paragraph_format.keep_with_next=True
    s.paragraph_format.space_before=Pt(7);s.paragraph_format.space_after=Pt(5)
D.styles['Title'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
D.styles['Caption'].font.size=Pt(10.5);D.styles['Caption'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
D.styles['Caption'].paragraph_format.line_spacing=1
def field(p,code,cached='1'):
    for tag in ['begin','instruction','separate','value','end']:
        r=p.add_run()
        if tag=='instruction':
            e=OxmlElement('w:instrText');e.set(qn('xml:space'),'preserve');e.text=' '+code+' ';r._r.append(e)
        elif tag=='value':r.text=cached
        else:
            e=OxmlElement('w:fldChar');e.set(qn('w:fldCharType'),tag);r._r.append(e)
def p(text='',style=None):return D.add_paragraph(text,style)
def h(text,level=2):return D.add_heading(text,level)
def page():D.add_page_break()
fig=0;tab=0
def caption(kind,text):
    global fig,tab
    if kind=='Hình':fig+=1;n=fig;seq='Figure'
    else:tab+=1;n=tab;seq='Table'
    pp=p('', 'Caption');pp.add_run(kind+' 5.');field(pp,'SEQ '+seq+' \\* ARABIC',str(n));pp.add_run('  '+text)
    return pp
def image(name,title,width=16.5):
    pp=p();pp.paragraph_format.space_after=Pt(0);pp.paragraph_format.keep_with_next=True;pp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    shape=pp.add_run().add_picture(str(AS/(name+'.png')),width=Cm(width));shape._inline.docPr.set('descr',title)
    caption('Hình',title)
def table(headers,rows,widths,title):
    cap=caption('Bảng',title);cap.paragraph_format.keep_with_next=True
    t=D.add_table(rows=1,cols=len(headers));t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for c,w in zip(t.columns,widths):c.width=Cm(w)
    for cell,text,w in zip(t.rows[0].cells,headers,widths):cell.width=Cm(w);cell.text=text
    for row in rows:
        cells=t.add_row().cells
        for c,text,w in zip(cells,row,widths):c.width=Cm(w);c.text=text
    for i,row in enumerate(t.rows):
        pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
        if i==0:pr.append(OxmlElement('w:tblHeader'))
        for c in row.cells:
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=c._tc.get_or_add_tcPr(); mar=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                a=OxmlElement('w:'+side);a.set(qn('w:w'),'70');a.set(qn('w:type'),'dxa');mar.append(a)
            cp.append(mar);shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E7ECE8' if i==0 else ('F5F7F5' if i%2 else 'FFFFFF'));cp.append(shade)
            for pp in c.paragraphs:
                pp.paragraph_format.space_after=Pt(2);pp.paragraph_format.space_before=Pt(1);pp.paragraph_format.line_spacing=1.02
                for r in pp.runs:r.font.size=Pt(11.5);r.bold=i==0
    p().paragraph_format.space_after=Pt(0)
def shot(text):
    pp=p('Ảnh UI cần bổ sung: '+text)
    for r in pp.runs:r.italic=True;r.font.size=Pt(11)
def code(text):
    pp=p('Đối chiếu mã nguồn: '+text)
    for r in pp.runs:r.font.size=Pt(10.5)
def link(p,text,url):
    a=OxmlElement('w:hyperlink');a.set(qn('r:id'),p.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True))
    r=OxmlElement('w:r');pr=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'1C5136');pr.append(color);r.append(pr);t=OxmlElement('w:t');t.text=text;r.append(t);a.append(r);p._p.append(a)

header=sec.header.paragraphs[0];header.text='RESCUEFARM  |  XÂY DỰNG ỨNG DỤNG';header.runs[0].font.size=Pt(9)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT;footer.add_run('Trang ');field(footer,'PAGE')
settings=D.settings.element;u=OxmlElement('w:updateFields');u.set(qn('w:val'),'true');settings.append(u)

# 1. Preserve the user's architecture figure with the high-resolution local export.
p('CHƯƠNG 5  XÂY DỰNG ỨNG DỤNG','Title')
h('5.1  Cấu trúc của ứng dụng',1)
p('RescueFarm được xây dựng bằng Java và XML Views trên Android, phục vụ Guest, Customer, Seller và Admin. Chương này trình bày cách tổ chức mã, kết nối Firebase và phối hợp dữ liệu giữa các trang. Phạm vi mô tả được đối chiếu với repository ngày 25/10/2026, tại commit 0d0f0c7; sự hiện diện của mã nguồn không đồng nghĩa dịch vụ production đã triển khai bản tương ứng.')
image('architecture-gitdiagram-print','Tổng quan các thành phần RescueFarm từ sơ đồ GitDiagram',15.7)
p('Hình 5.1 tổ chức hệ thống thành ứng dụng Android, tầng truy cập dữ liệu và backend tin cậy. Ở Android, Fragment hiển thị nội dung, ViewModel giữ trạng thái màn hình và Repository phối hợp Firebase với bộ nhớ cục bộ. Các tác vụ như checkout, kiểm tra tồn kho và xác minh thanh toán được xử lý tập trung tại server để tránh tin cậy dữ liệu do client tự gửi.')
pp=p('Nguồn hình: ');link(pp,'GitDiagram của repository RescueFarm','https://gitdiagram.com/laithanhdat-glitch/nhom_7_mobileapp');pp.runs[0].font.size=Pt(10)

# 2. Concrete package mapping rather than another theory chapter.
page();h('5.1  Cấu trúc của ứng dụng tiếp theo',1)
p('Mã nguồn được chia theo trách nhiệm, trong khi các màn hình được nhóm tiếp theo nghiệp vụ. Cấu trúc này giúp theo dõi một thao tác từ giao diện đến nơi lưu dữ liệu, đồng thời cho phép thay nguồn truy cập hoặc kiểm thử logic mà không phải viết lại toàn bộ màn hình.')
table(['Thành phần','Vai trò trong triển khai','Ví dụ trong repository'],[
 ('ui/','Fragment, ViewModel, adapter và trạng thái theo màn.','HomeViewModel; CampaignDetailFragment; CartViewModel.'),
 ('domain/','Model, enum và điều kiện nghiệp vụ độc lập giao diện.','Product, ProductBatch, RescueCampaign, Comment.'),
 ('data/repository/','Hợp đồng đọc ghi; factory chọn nguồn theo phiên.','ProductRepository; CartRepositoryFactory.'),
 ('data/remote/firebase/','Query, listener, ánh xạ document và xử lý quyền.','FirebaseProductRepository; SnapshotLiveData.'),
 ('data/local/','Room database, entity, DAO, mapper; lưu tùy chọn cục bộ.','RescueFarmDatabase phiên bản 11; HomeLocationStore.'),
 ('service/','HTTP backend, GPS, media, giá/kho và đồng bộ nền.','RenderBackendClient; MobileLocationService; ReadCacheSyncWorker.'),
 ('res/','Layout XML, biến thể màn hình, drawable và navigation.','layout; layout-land; navigation graph.'),
 ('functions/','Express trên Render; nghiệp vụ cần kiểm tra phía server.','commerce.js; sepay.js; rescue-fund.js.'),
 ('cloudinary-worker/','Xác thực và ký thao tác media Cloudinary.','src/index.js; không giữ API secret trong APK.')
],[3.2,6.1,7.4],'Ánh xạ cấu trúc mã nguồn và trách nhiệm')
p('Không phải mọi Fragment đều bắt buộc đi qua cùng một lớp ViewModel: một số màn gọi Repository hoặc HTTP client trực tiếp. MVVM và Repository là cách tổ chức chủ đạo, còn đường đi cụ thể phải được đọc theo từng chức năng. Thư mục functions chứa ứng dụng Express triển khai trên Render, không phải bằng chứng hệ thống chạy Firebase Cloud Functions.')
p('Dữ liệu công khai đã xem như danh mục, sản phẩm, lô, chiến dịch và banner có các bảng cache Room. Đơn hàng cache được gắn với người dùng; giỏ Guest lưu theo phiên khách. TTL chính sách là 30 phút cho catalog, 10 phút cho chiến dịch/đơn và 1 giờ cho banner; mốc hard expiry là 7 ngày. Cache hỗ trợ đọc lại dữ liệu đủ điều kiện, không thay thế kiểm tra giá, quyền và tồn kho online.')
p('Gradle quản lý cấu hình build và thư viện qua app/build.gradle.kts cùng version catalog. Mã Java biên dịch ở mức tương thích Java 11; minSdk 29 và compile/targetSdk 37. Bộ kiểm thử được tách thành JVM, instrumentation và Firestore Emulator để kiểm tra đúng tầng, thay vì dùng kết quả build thay cho kiểm thử chức năng.')
code('app/build.gradle.kts; gradle/libs.versions.toml; data/local/database/RescueFarmDatabase.java; service/sync/CacheSyncPolicy.java.')

# 3. Firebase integration.
page();h('5.2  Kết nối với Firebase',1)
p('Ứng dụng đăng ký với Firebase theo applicationId com.example.nhom_7_mobileapp. Tệp cấu hình google-services.json được đặt trong module app và xử lý bởi plugin Google Services. Firebase BoM quản lý tương thích thư viện Auth, Firestore, Messaging và App Check; RescueFarmApplication khởi tạo Firebase và cài provider App Check theo biến thể build.')
image('firebase-connection','Hai đường truy cập Firebase qua SDK client và backend tin cậy')
table(['Dịch vụ hoặc cấu hình','Cách sử dụng trong RescueFarm'],[
 ('Authentication','Email/mật khẩu; Google Credential Manager lấy Google ID token rồi đổi thành Firebase credential. UID liên kết hồ sơ users.'),
 ('Firestore và Rules','Lưu collection/document; Rules kiểm vai trò, trạng thái và quyền sở hữu của thao tác SDK. Index phục vụ query hiện có.'),
 ('Render và Admin SDK','Xác thực ID token và kiểm nghiệp vụ phía server. Admin SDK không tự chịu giới hạn Rules của client.'),
 ('Media và thông báo','Ảnh lưu Cloudinary qua Worker ký upload; Firestore giữ URL/metadata. FCM hỗ trợ nhận thông báo; không bảo đảm đến tức thời.')
],[4.4,12.3],'Cấu hình kết nối và phạm vi từng dịch vụ')
p('Khi đăng nhập thành công, ứng dụng đọc hồ sơ users/{uid} để áp dụng vai trò và trạng thái tài khoản. Ẩn nút trên UI chỉ hỗ trợ trải nghiệm; quyền dữ liệu vẫn do Rules hoặc backend kiểm tra. Service account và khóa ký media chỉ nằm phía server, không nhúng vào APK. Room là SQLite của ứng dụng, độc lập với cache nội bộ Firestore SDK; URL ảnh trong Room không chứng minh ảnh đã có trên đĩa.')
shot('Firebase Console phần đăng ký ứng dụng Android và danh sách collection demo; che project nhạy cảm. Có thể thêm cặp màn đăng nhập email/Google khi kiểm chứng provider thật, không dùng ảnh này thay cho bằng chứng Rules.')
code('RescueFarmApplication; FirebaseAuthRepository.signInWithGoogleIdToken; firestore.rules; firestore.indexes.json. Môi trường reportDemo nối Auth/Firestore Emulator và HTTP local riêng.')

# 4. Home and search.
page();h('5.3  Xây dựng nội dung các trang',1)
h('a  Trang chủ và trang tìm kiếm')
p('Trang chủ tổng hợp sản phẩm, chiến dịch, banner và các nhóm giải cứu gần người xem. HomeViewModel theo dõi các nguồn Repository, kết hợp trạng thái tải, vị trí và bộ lọc thành HomeViewState. Thẻ chiến dịch lấy ảnh đại diện từ ảnh chiến dịch hiện có; chạm thẻ điều hướng theo ID đến trang chi tiết. Vị trí tìm kiếm được lưu cục bộ theo Guest hoặc UID; dữ liệu vị trí của tài khoản được xóa khi đăng xuất.')
image('home-search','Luồng đọc dữ liệu Home và xử lý tìm kiếm trên tập đã tải')
p('Thao tác với Firebase bắt đầu bằng query sản phẩm ACTIVE, chiến dịch ACTIVE, danh mục và banner đủ điều kiện. Trang công khai ban đầu hiển thị tối đa 12 sản phẩm/chiến dịch; query lấy thêm một document để xác định còn trang. Repository ánh xạ snapshot thành model, lưu cache qua DAO rồi phát LiveData. Khi màn không còn observer, SnapshotLiveData tháo listener theo vòng đời.')
p('Trang tìm kiếm dùng DiscoveryEngine xử lý từ khóa, danh mục, khoảng giá, địa phương và cách sắp xếp trên dữ liệu đã tải. Khoảng cách được tính từ tọa độ, không phải truy vấn tìm kiếm toàn văn của Firestore. Người dùng tải thêm để mở rộng tập kết quả; không có cơ sở khẳng định kết quả ban đầu đã bao phủ toàn bộ database. Khi mất mạng, UI có thể đọc cache còn hiệu lực nhưng phải thể hiện độ mới của dữ liệu.')
table(['Đầu vào','Xử lý và đầu ra'],[
 ('Từ khóa và bộ lọc','HomeSearchInput chuẩn hóa thao tác nhập; DiscoveryQuery/DiscoveryEngine tạo danh sách phù hợp.'),
 ('GPS hoặc vị trí chọn','HomeLocationStore giữ lựa chọn theo phiên; nhóm gần bạn có khoảng cách khi đủ tọa độ.'),
 ('Làm mới hoặc tải thêm','Repository cập nhật query; lỗi mạng giữ nội dung đã tải thay vì giả danh sách mới.')
],[4.4,12.3],'Đầu vào và kết quả của nhóm khám phá')
shot('Một ảnh Home có thẻ chiến dịch/ảnh thật và bản đồ nhỏ; một ảnh tìm kiếm có từ khóa, bộ lọc và kết quả. Dùng cùng dữ liệu để chụp dọc/ngang; thêm ảnh offline có nhãn dữ liệu đã lưu nếu đã kiểm chứng.')
code('ui/home/HomeViewModel.java; DiscoveryEngine.java; data/remote/firebase/FirebaseProductRepository.java; FirebaseCampaignRepository.java; SnapshotLiveData.java.')

# 5. Product/cart/checkout and order data.
page();h('b  Trang chi tiết sản phẩm và c  Trang giỏ hàng')
p('Chi tiết sản phẩm hiển thị ảnh, mô tả, nguồn gốc và các lô có thể mua. Product mô tả mặt hàng; ProductBatch chứa hạn dùng, giá và tồn của từng lô. Người mua chọn lô và số lượng trước khi thêm giỏ. ProductViewModel phối hợp ProductRepository, kiểm tra tính khả dụng và giữ lựa chọn UI; không dùng một ngày hết hạn chung của Product để thay cho hạn lô.')
p('CartRepositoryFactory chọn GuestCartRepository dùng Room hoặc FirebaseCartRepository dùng carts/{customerId}. Mỗi dòng giữ productId, batchId, sellerId, số lượng và báo giá tham khảo. Thêm/sửa giỏ được tái kiểm tra bằng báo giá mới; giỏ chưa đồng nghĩa đã giữ hàng. Xóa giỏ chỉ xóa dữ liệu giỏ, không xóa sản phẩm/lô hay lịch sử đơn.')
image('commerce','Luồng từ chọn lô và giỏ hàng đến checkout giữ kho')
h('d  Trang thanh toán và theo dõi đơn hàng')
p('Checkout gửi thông tin nhận hàng, phương thức giao/thu tiền và requestId đến /checkoutCustomer hoặc /checkoutGuest. Backend tính lại giá và phí, kiểm tra người bán, hạn lô, phân bổ chiến dịch và tồn khả dụng rồi ghi đơn, giữ kho và checkoutReceipts trong transaction. Nếu báo giá thay đổi, UI yêu cầu người mua xem và xác nhận lại; retry cùng yêu cầu không được tạo thêm đơn hoặc giữ kho lần hai.')
p('Đơn dùng snapshot mặt hàng/giá để giữ lịch sử khi dữ liệu gốc thay đổi. Trang đơn đọc trạng thái theo quyền của chủ đơn hoặc Seller liên quan; các chuyển trạng thái quan trọng qua API. Với QR, webhook SePay hợp lệ xác nhận tiền vào, không có thao tác tự sửa số dư để đánh dấu đã thanh toán. Với COD, trạng thái thu tiền đi theo vòng đời đơn. Checkout bắt buộc online; cache chỉ hỗ trợ tra cứu.')
shot('Chi tiết có ảnh và bộ chọn lô/số lượng; giỏ có tăng/giảm/xóa; checkout có phương thức nhận/thanh toán; đơn sau tạo. Nếu minh họa QR, ghi rõ demo và chỉ chụp trạng thái đã trả khi có webhook emulator hợp lệ.')
code('ProductViewModel; CartViewModel; CartRepositoryFactory; FirebaseOrderRepository.checkout; functions/commerce.js; functions/sepay.js.')

# 6. Campaign, location and community.
page();h('e  Trang chiến dịch và bản đồ điểm giải cứu')
p('Chi tiết chiến dịch trình bày mục tiêu, lô liên quan, địa điểm và tiến độ. FIXED_POINT dùng tọa độ điểm cố định; MOBILE_POINT dùng tọa độ hiện tại do Seller sở hữu chiến dịch chủ động chia sẻ. Map preview trên Home mở bản đồ lớn, marker mở thông tin chiến dịch. Leaflet trong WebView hiển thị bản đồ; mở ứng dụng chỉ đường ngoài không đồng nghĩa app tích hợp Google Maps SDK.')
image('campaign-map','Luồng GPS chia sẻ vị trí và bình luận có ảnh của chiến dịch')
p('MobileLocationService là foreground service loại location. Seller phải cấp quyền và bật chia sẻ; vị trí hợp lệ được giới hạn tần suất theo LocationPublishPolicy trước khi Repository ghi currentLatitude/currentLongitude, locationUpdatedAt và cờ chia sẻ. Backend rescue-map.js dùng một listener campaigns ACTIVE dùng chung trong mỗi process, lọc các trường công khai rồi đẩy SSE tới người xem. Marker thay đổi theo dữ liệu mới nhận, không polling Firestore mỗi giây. Vị trí mobile thiếu timestamp, quá 10 phút hoặc ngừng chia sẻ không được coi là điểm đang hoạt động.')
h('Bình luận và tương tác cộng đồng',3)
p('Customer có thể gửi text 1–500 ký tự kèm tối đa ba ảnh. Ảnh được upload Cloudinary bằng chữ ký Worker; /campaigns/comments/add kiểm quyền, chủ sở hữu ảnh và lịch chiến dịch rồi ghi subcollection comments bằng transaction với commentId ổn định. Backend lấy tên/avatar từ hồ sơ, không tin thông tin tác giả client tự khai. Listener hiển thị 100 bình luận mới nhất với tên, avatar, thời gian, nội dung và ảnh; chạm ảnh để xem lớn.')
p('Bình luận ẩn danh dùng tên và avatar mặc định trên UI, UID vẫn tồn tại phục vụ sở hữu/kiểm duyệt. Bình luận cũ thiếu thông tin tác giả được bổ sung qua endpoint trình bày có cache 60 giây. Reaction/report được giữ theo quyền hiện có. Số reads vẫn phụ thuộc số document và thay đổi; listener chia sẻ cùng cache chỉ giảm tải, không bảo đảm miễn phí vô hạn.')
shot('Cùng một chiến dịch MOBILE_POINT: màn Seller bật chia sẻ và bản đồ người xem tại hai tọa độ demo; thêm cặp bình luận thường/ẩn danh có ảnh. Ghi giờ cập nhật, không dùng ảnh cũ làm bằng chứng xe đang di chuyển.')
code('MobileLocationService; LocationPublishPolicy; FirebaseCampaignRepository.updateMobileLocation; functions/rescue-map.js; campaign-comments.js; CampaignDetailFragment.')

# 7. Seller and admin mobile.
page();h('f  Trang quản lý của Seller và g  Trang quản trị')
p('Seller quản lý sản phẩm, lô và chiến dịch qua các màn Android theo tài khoản được duyệt. Form kiểm tra trường bắt buộc và gửi dữ liệu qua Repository. Product giữ thông tin chung và ảnh; Batch giữ số lượng, giá, hạn dùng và liên kết Product. Seller phải thuộc đúng chủ sở hữu; lô hết hạn không được tiếp tục bán bằng cách thay trạng thái hiển thị. Các cập nhật không được làm mất lịch sử giao dịch đã phát sinh.')
image('seller-admin','Phối hợp dữ liệu giữa Seller và Admin trên ứng dụng Android')
table(['Nhóm tác vụ','Thao tác dữ liệu và kiểm soát'],[
 ('Sản phẩm và ảnh','Repository ghi products theo quyền Seller; Worker ký upload media, document giữ URL/public ID.'),
 ('Quản lý lô','Ghi productBatches theo sản phẩm thuộc Seller; kiểm hạn và tính nhất quán tồn; không coi lượng đang giữ là hàng còn bán.'),
 ('Tạo chiến dịch','Chọn lô/mục tiêu/lịch và chế độ điểm; ghi campaigns cùng liên kết phân bổ theo nghiệp vụ hiện có.'),
 ('Kiểm duyệt mobile','Admin xem yêu cầu Seller và nội dung; duyệt/từ chối, ghi trạng thái và thông tin kiểm duyệt theo Rules/API.'),
 ('Đơn và vận hành','Seller xử lý đơn của mình; Admin xem và đối soát theo quyền. Không sửa đơn hoặc thanh toán trực tiếp từ client.')
],[4.3,12.4],'Ánh xạ tác vụ quản lý vào dữ liệu Firebase')
p('Firestore Rules kiểm tra vai trò, trạng thái tài khoản và sellerId/UID; việc Seller nhìn thấy một nút không đủ tạo quyền ghi. Với API Admin SDK, backend phải tự xác thực và kiểm tra trạng thái vì Rules của client không được áp dụng tự động. Dữ liệu đơn/kho, thanh toán và hũ cần dùng đúng endpoint, không cho UI gán tùy ý các bộ đếm.')
p('Khi Admin duyệt, listener/refresh của màn liên quan đưa trạng thái mới về giao diện. Đường dẫn ID giữ sự liên kết giữa danh sách và chi tiết, tránh thay document ID bằng tên hiển thị. Banner cần liên kết campaignId hợp lệ mới điều hướng đúng; một ảnh banner tồn tại không chứng minh cấu hình đích đã đầy đủ.')
shot('Form Seller có sản phẩm và lô thật trong demo; form chiến dịch với chế độ điểm và lịch; Admin mobile mở một nội dung chờ duyệt rồi xem kết quả. Che dữ liệu cá nhân; không chụp web Console thay cho UI Admin.')
code('ui/product; ui/campaign; ui/admin; FirebaseProductRepository; FirebaseCampaignRepository; firestore.rules; functions/campaign-media.js.')

# 8. Fund and operational boundaries.
page();h('h  Trang Hũ giải cứu và chuyến thiện nguyện')
p('Màn Hũ hiển thị số tiền có thể phân bổ, đang giữ, đã chi, nơi đón nhận và sổ biến động. Người ủng hộ dùng QR/mã theo dõi; webhook SePay hợp lệ tự ghi nhận tiền vào và chống trùng. Seller theo dõi chuyến, nhận chuyến và gửi ảnh giao; Admin xác minh rồi lưu thông tin chi. Danh sách chuyến được mở ở trang riêng, vẫn nằm trong ứng dụng mobile.')
image('rescue-fund','Luồng tiền và hàng trong Hũ giải cứu từ ủng hộ đến ghi chi')
p('Tại thời điểm phân bổ t, backend chọn lô có expiryDate > t + 24 giờ + 1 phút và ≤ t + 48 giờ. Lô gắn chiến dịch còn cần chiến dịch ACTIVE, đúng Seller/đơn vị và kết thúc sau t + 24 giờ + 1 phút. Cửa sổ không phát sinh mua thành công 12/24 giờ là một điều kiện riêng. Transaction kiểm ngân sách, tiền/tồn khả dụng và chống phân bổ lặp; đồng thời giữ tiền, giữ hàng, tạo chuyến và ledger.')
p('Chuyến mới có deadline = allocatedAt + 24 giờ, deliveryWindowHours = 24 và allocationPolicyVersion = 2. Chuỗi trạng thái là RESERVED → ACCEPTED → DELIVERED → VERIFIED → PAID. Seller phải gửi bằng chứng trước hạn; nhận chuyến không cộng lại 24 giờ. Worker hủy RESERVED/ACCEPTED quá hạn và giải phóng đúng một lần. DELIVERED đúng hạn vẫn được Admin duyệt muộn; hủy chỉ trả khoản đang giữ về hũ, không hoàn ngân hàng cho người ủng hộ.')
p('VERIFY chốt kho nhưng tiền vẫn reserved. Admin chuyển tiền ngoài app rồi PAY lưu số tiền, mã giao dịch và ảnh chứng từ; lúc đó mới chuyển reserved sang spent. Số dư tuân theo available + reserved + spent = tổng ủng hộ xác nhận + điều chỉnh đối soát, khởi đầu bằng 0; không đồng nhất với số dư tài khoản ngân hàng dùng chung. RescueFundTime hiển thị giờ VN UTC+7; không đổi timestamp để đổi múi giờ.')
shot('Một mã chuyến demo xuyên suốt: tổng quan Hũ → Seller nhận/gửi ảnh có deadline → Admin VERIFIED/PAID và ledger. Thêm một chuyến quá hạn riêng để thấy RELEASE. Không tự gán status hoặc sửa số dư để tạo ảnh.')
code('ui/fund/RescueFundFragment.java; RescueFundTime.java; functions/rescue-fund.js; docs/rescue-fund.md. Mô tả theo source; ảnh UI/provider/production mới cần được kiểm chứng riêng.')

# 9. References only, as requested. No bracketed in-text references.
page();h('TÀI LIỆU THAM KHẢO',1)
p('Các tài liệu dưới đây hỗ trợ giải thích kỹ thuật. Hành vi riêng của RescueFarm được đối chiếu mã nguồn, không suy ra từ tài liệu SDK. Ngày truy cập: 25/10/2026.')
refs=[
 ('Google Firebase','Add Firebase to your Android project','https://firebase.google.com/docs/android/setup','Cấu hình ứng dụng Android, plugin Google Services và thư viện Firebase.'),
 ('Google Firebase','Authenticate with Google on Android','https://firebase.google.com/docs/auth/android/google-signin','Credential Manager và chuyển Google credential sang phiên Firebase Authentication.'),
 ('Google Firebase','Get realtime updates with Cloud Firestore','https://firebase.google.com/docs/firestore/query-data/listen','Snapshot listener, dữ liệu ban đầu và cập nhật sau đó.'),
 ('Google Firebase','Transactions and batched writes','https://firebase.google.com/docs/firestore/manage-data/transactions','Thao tác nguyên tử và xử lý cạnh tranh dữ liệu bằng transaction.'),
 ('Google Firebase','Security Rules and Firebase Authentication','https://firebase.google.com/docs/rules/rules-and-auth','Kết hợp xác thực và kiểm soát truy cập ở tầng dữ liệu.'),
 ('Android Developers','Save data in a local database using Room','https://developer.android.com/training/data-storage/room','Room, entity/DAO và SQLite phục vụ dữ liệu cục bộ.'),
 ('Google Firebase','Connect your app to the Cloud Firestore Emulator','https://firebase.google.com/docs/emulator-suite/connect_firestore','Cô lập kiểm thử với Firestore Emulator và project demo.'),
 ('GitDiagram','Sơ đồ repository RescueFarm do người dùng cung cấp','https://gitdiagram.com/laithanhdat-glitch/nhom_7_mobileapp','Nguồn hình tổng quan 5.1; dùng bản xuất độ phân giải cao đã có trong repository.')]
for i,(org,title,url,desc) in enumerate(refs,1):
    pp=p(f'{i}. {org}. {title}.');pp.paragraph_format.keep_with_next=True
    pp=p();link(pp,url,url);pp.paragraph_format.keep_with_next=True
    pp=p(desc);pp.paragraph_format.space_after=Pt(10)
    for r in pp.runs:r.font.size=Pt(11.5)

doc=OUT/'Phan-5-Xay-dung-ung-dung.docx';D.save(doc)
evidence=dict(date='2026-10-25',baseCommit=subprocess.check_output(['git','rev-parse','HEAD'],cwd=ROOT,text=True).strip(),
 sourcePdf='docs/report/Phần của Đạt.pdf',sourceImage='docs/diagram (2).png',sourceImagePixels=Image.open(AS/'architecture-gitdiagram-hd.png').size,
 figureCount=fig,tableCount=tab,plannedPages={'content':8,'references':1},
 limits=['No new UI screenshots; precise shot notes supplied','No tests or production deploy claimed in this document task','GitDiagram web viewer inaccessible via web fetch; provided local HD export used'],
 references=[dict(organization=o,title=t,url=u) for o,t,u,_ in refs])
(OUT/'source/evidence.json').write_text(json.dumps(evidence,ensure_ascii=False,indent=2),encoding='utf-8')
print('CREATED',doc,'FIGURES',fig,'TABLES',tab)
