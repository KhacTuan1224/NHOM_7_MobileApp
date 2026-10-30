"""Source excerpts shared by the chapter-five document and report builders."""
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

EXCERPTS = {
'architecture-gitdiagram-print': dict(title='Factory tạo Repository và phối hợp Room với Firebase',
 source='data/repository/ProductRepositoryFactory.java, create(), dòng 13–20',
 code='''RescueFarmDatabase database = RescueFarmDatabase.getInstance(context);
if (!FirebaseConfiguration.isConfigured(context)) {
    return new OfflineProductRepository(
        database.catalogCacheDao(), ...);
}
return new FirebaseProductRepository(database, ...,
        new AndroidNetworkStatusProvider(context));''',
 explanation='Factory khởi tạo Room và chọn Repository phù hợp với cấu hình Firebase. Khi có Firebase, Repository nhận cả database cục bộ lẫn bộ theo dõi mạng để phối hợp nguồn đọc. Điều kiện isConfigured kiểm tra cấu hình dịch vụ, không tự đồng nghĩa thiết bị đang có mạng.'),
'firebase-connection': dict(title='Firebase Authentication tiếp nhận Google credential',
 source='data/remote/firebase/FirebaseAuthRepository.java, signInWithGoogleIdToken(), từ dòng 104',
 code='''AuthCredential credential = GoogleAuthProvider.getCredential(idToken, null);
firebaseAuth.signInWithCredential(credential)
    .addOnSuccessListener(result -> {
        FirebaseUser firebaseUser = result.getUser();
        ...
        loadOrCreateGoogleProfile(firebaseUser, callback);
    })
    .addOnFailureListener(exception -> notifyFailure(exception, callback));''',
 explanation='Google ID token được đổi thành Firebase credential để tạo phiên xác thực. Sau khi thành công, Repository đọc hoặc tạo hồ sơ ứng dụng để lấy vai trò và trạng thái; nhánh lỗi trả thông tin qua callback. Xác thực danh tính không thay thế kiểm tra quyền tại Rules hoặc backend.'),
'home-search': dict(title='HomeViewModel nhận dữ liệu và phát kết quả tìm kiếm',
 source='ui/home/HomeViewModel.java, từ dòng 97 và publishSearch() từ dòng 431',
 code='''homeState.addSource(productRepository.observePublicProductPage(), values -> {
    products = safe(values);
    ...
    publish();
});
...
private void publishSearch() {
    searchResults.postValue(discoveryEngine.search(
        products, campaigns, productBatches, sellerCoordinates, query,
        latitude, longitude, new Date()));
}''',
 explanation='Repository phát danh sách sản phẩm; ViewModel cập nhật dữ liệu rồi dựng lại trạng thái. Khi tìm kiếm, DiscoveryEngine nhận tập dữ liệu đã tải, bộ lọc và tọa độ, sau đó LiveData phát kết quả cho màn hình. Hai đoạn thuộc các vị trí khác nhau trong cùng file; dấu ba chấm giữa chúng lược phần mã không liên quan.'),
'commerce': dict(title='Chọn nguồn lưu giỏ theo phiên người dùng',
 source='data/repository/CartRepositoryFactory.java, create(), từ dòng 13',
 code='''AuthRepository auth = AuthRepositoryFactory.create(context);
if (auth.isAuthenticated() && auth.getCurrentUserId() != null) {
    return new FirebaseCartRepository(auth.getCurrentUserId(), ...);
}
...
return new GuestCartRepository(
    RescueFarmDatabase.getInstance(context).guestCartDao(), ...);''',
 explanation='Factory chọn giỏ tài khoản qua FirebaseCartRepository hoặc giỏ Guest qua DAO của Room. Định danh Guest và các tham số executor được lược để tập trung vào việc chọn nguồn dữ liệu. Giỏ lưu lựa chọn mua, không phải một lần giữ tồn kho.'),
'checkout': dict(title='Tạo đơn và chuyển lượng khả dụng sang đang giữ',
 source='functions/commerce.js, writeOrder(), dòng 222 và 243; checkout() mở transaction ở dòng 295',
 code='''tx.create(orderRef, order);
...
tx.update(row.batchRef, {
  availableQuantity: batch.availableQuantity - item.quantity,
  reservedQuantity: batch.reservedQuantity + item.quantity,
  status: 'RESERVED_PARTIAL',
  ...
});''',
 explanation='Đoạn ghi nằm trong writeOrder(tx, ...), được checkout() gọi bên trong db.runTransaction. Sau khi kiểm tra lại dữ liệu, backend tạo đơn và chuyển lượng khách đặt từ khả dụng sang đang giữ trong cùng giao dịch. Nếu transaction thất bại thì không commit một phần; lượng đang giữ chưa được coi là đã bán hoặc đã giao.'),
'campaign-map': dict(title='Foreground service gửi vị trí điểm giải cứu di động',
 source='service/location/MobileLocationService.java, publish(), từ dòng 78; lời gọi ở dòng 86',
 code='''private void publish(Location fix) {
    ...
    if (System.currentTimeMillis() >= endAt) {
        stopSharing(); return;
    }
    ...
    repository.updateMobileLocation(
        campaignId, fix.getLatitude(), fix.getLongitude(), true,
        new CampaignRepository.CampaignCallback() {
            ...
        });
}''',
 explanation='Service kiểm tra điều kiện chia sẻ rồi gửi tọa độ hiện tại của Campaign qua Repository. Các kiểm tra chất lượng, tuổi fix và tiết chế tần suất được lược trong trích đoạn. Backend bản đồ phát điểm đủ điều kiện qua SSE; vị trí này thuộc điểm giải cứu, không tự trở thành GPS vận chuyển của từng đơn.'),
'comments': dict(title='Bình luận lấy danh tính từ hồ sơ và lưu qua transaction',
 source='functions/campaign-comments.js, endpoint /campaigns/comments/add, từ dòng 32 và 42',
 code='''const author = identity(profile.data(), b.anonymous === true),
      ref = db.collection('campaigns').doc(b.campaignId)
              .collection('comments').doc(b.commentId);
...
tx.set(ref, {
  id: b.commentId, campaignId: b.campaignId, userId: caller.uid,
  status: 'ACTIVE', content: b.content.trim(),
  senderName: author.name, senderAvatarUrl: author.avatarUrl,
  anonymous: author.anonymous, imageUrls: images,
  ...
});''',
 explanation='Backend lấy tên và avatar từ hồ sơ đã đọc, áp dụng lựa chọn ẩn danh rồi ghi bình luận theo commentId. UID người gửi vẫn được giữ để kiểm tra sở hữu và kiểm duyệt. Các bước kiểm tra độ dài nội dung, tối đa ba ảnh, quyền và chống ghi lặp nằm trước hoặc bao quanh đoạn ghi được trích.'),
'seller-admin': dict(title='Kiểm duyệt Campaign theo chuyển trạng thái được phép',
 source='service/admin/AdminTransitionPolicy.java, allowedTargets(), từ dòng 25',
 code='''case CAMPAIGNS:
    if (status.equals("PENDING_APPROVAL"))
        return list("ACTIVE", "REJECTED");
    return status.equals("ACTIVE") ? list("STOPPED") : none();''',
 explanation='Policy xác định trạng thái đích hợp lệ: Campaign chờ duyệt có thể được duyệt hoặc từ chối; Campaign đang hoạt động có thể bị dừng. Đây là bước kiểm tra quy tắc, không phải lệnh ghi Firestore. Repository/API còn phải xác minh quyền Admin và yêu cầu lý do ở hành động tương ứng.'),
'rescue-fund': dict(title='Hũ phân bổ chuyến và giữ đồng thời tiền với hàng',
 source='functions/rescue-fund.js, allocation(), transaction từ dòng 296, phần ghi từ dòng 337',
 code='''tx.create(jobRef, job);
tx.update(batchRef, {
  availableQuantity: round(batch.availableQuantity - quantity),
  reservedQuantity: round(batch.reservedQuantity + quantity),
  ...
});
tx.update(fundRef, {
  available: config.available - amount,
  reserved: (config.reserved || 0) + amount,
  ...
});''',
 explanation='Sau khi kiểm tra lô và ngân sách, transaction tạo chuyến rồi giảm lượng hàng/tiền khả dụng, tăng phần đang giữ tương ứng. quantity là lượng hàng phân bổ và amount là ngân sách giữ cho chuyến. Bước này chưa trả tiền Seller; VERIFY và PAY thực hiện sau theo bằng chứng và khoản chuyển ngoài ứng dụng.'),
}

def add_excerpt(document, key, number):
    item = EXCERPTS[key]
    p=document.add_paragraph()
    p.paragraph_format.keep_with_next=True
    p.paragraph_format.space_before=Pt(6)
    r=p.add_run(f'Mã nguồn 5.{number}. {item["title"]}');r.bold=True;r.font.size=Pt(11)
    p=document.add_paragraph()
    p.alignment=WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.line_spacing=1.0
    p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(4)
    p.paragraph_format.keep_together=True;p.paragraph_format.keep_with_next=True
    p.paragraph_format.left_indent=Pt(8);p.paragraph_format.right_indent=Pt(5)
    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'F3F5F6');p._p.get_or_add_pPr().append(shade)
    r=p.add_run(item['code']);r.font.name='Consolas';r.font.size=Pt(9)
    r.font.color.rgb=RGBColor(25,35,45)
    r._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Consolas')
    p=document.add_paragraph('Nguồn: '+item['source']+'.')
    p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=1.0
    p.paragraph_format.keep_with_next=True
    for r in p.runs:r.italic=True;r.font.size=Pt(9)
    p=document.add_paragraph(item['explanation'])
    p.paragraph_format.space_after=Pt(6)
    p.paragraph_format.keep_together=True

def markdown_excerpt(key, number):
    item=EXCERPTS[key]
    language='javascript' if item['source'].startswith('functions/') else 'java'
    return (f'**Mã nguồn 5.{number}. {item["title"]}**\n\n'
            f'```{language}\n{item["code"]}\n```\n\n'
            f'Nguồn: `{item["source"]}`.\n\n{item["explanation"]}\n')
