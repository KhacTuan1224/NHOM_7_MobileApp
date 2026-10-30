"""Embed original pictures at their original display size and keep code blocks together."""
from pathlib import Path
from io import BytesIO
import re
from lxml import html
from PIL import Image
from docx import Document
from docx.shared import Inches,Pt
from docx.text.paragraph import Paragraph
from docx.text.run import Run

PART=Path(__file__).resolve().parents[1]
source=PART/'source/chapter5-preserved.html'
root=html.fromstring(source.read_text(encoding='utf-8'))
images=root.xpath('//body//img')
target=PART/'Phan-5-Xay-dung-ung-dung.docx'
doc=Document(target)
shapes=list(doc.inline_shapes)
assert len(shapes)==len(images),(len(shapes),len(images))
for shape,img in zip(shapes,images):
    data=(source.parent/img.get('src')).resolve().read_bytes()
    with Image.open(BytesIO(data)) as picture:pw,ph=picture.size
    style=img.get('style','')
    w=re.search(r'(?:^|;)\s*width:\s*([0-9.]+)px',style)
    h=re.search(r'(?:^|;)\s*height:\s*([0-9.]+)px',style)
    width=float(w.group(1)) if w else pw
    height=float(h.group(1)) if h else ph
    old=shape._inline.getparent();run_element=old.getparent()
    paragraph_element=run_element.getparent()
    while not paragraph_element.tag.endswith('}p'):paragraph_element=paragraph_element.getparent()
    paragraph=Paragraph(paragraph_element,doc._body)
    run=Run(run_element,paragraph)
    new=run.add_picture(BytesIO(data),width=Inches(width/96),height=Inches(height/96))._inline.getparent()
    run_element.replace(old,new)
    if w:paragraph.paragraph_format.keep_with_next=True

inside=False;next_explanation=False
for p in doc.paragraphs:
    t=p.text.strip()
    if t.startswith('Mã nguồn 5.'):
        inside=True;p.paragraph_format.keep_with_next=True
    elif inside:
        p.paragraph_format.keep_with_next=True
        p.paragraph_format.keep_together=True
        if t.startswith('Nguồn:'):
            inside=False;next_explanation=True
        else:
            p.paragraph_format.space_after=Pt(0)
            p.paragraph_format.space_before=Pt(0)
            p.paragraph_format.line_spacing=1.0
    elif next_explanation and t:
        p.paragraph_format.keep_together=True;next_explanation=False
    if t.startswith(('Hình 5.1','Hình 5.2.')):p.paragraph_format.keep_with_next=True

# Remove unused external image links left by Word's HTML importer.
used={v for e in doc.element.body.iter() for k,v in e.attrib.items() if k.startswith('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}')}
for rid,rel in list(doc.part.rels.items()):
    if rel.reltype.endswith('/image') and rid not in used:doc.part.drop_rel(rid)
doc.save(target)
print('Embedded all original figures at source display sizes; code and one-paragraph explanations kept together.')
